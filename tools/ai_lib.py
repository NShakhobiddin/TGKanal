# -*- coding: utf-8 -*-
"""«AI darslar» kanali (channels/ai) uchun haftalik generator kutubxonasi.

Haftalik skript (tools/ai_week_NN.py) shu kutubxonadan foydalanadi:

    from ai_lib import Week, pre
    W = Week(1, start="2026-10-12", theme="Excel faylni kod bilan tahlil qildirish")
    f = W.file("murojaatlar_2025.xlsx", build_fn)            # mashq fayli (faqat yo'q bo'lsa yaratiladi)
    W.post("sesh", rubric="#dars", title="…", head="…", body="…", new="…", take="…",
           card={"kind": "chain", "steps": [...]}, attach=[{"file": f, "caption": "…"}])
    W.build()

Nima qiladi:
  • post JSON larini channels/ai/posts/ ga yozadi (faqat yangi fayllar; --rework — qoralamalarni qayta yozadi);
  • har post uchun kartochka rasmini o'zi chizadi (Pillow, 1080×1350) -> channels/ai/images/<id>-auto.png;
  • har postda «Bugungi yangi bilim» va «Olib keting» bo'lishini talab qiladi (kanal qoidasi).

Buyruq satri:  --start=2026-10-19  boshqa haftaga ko'chirish (shu haftaning eski qoralamalari o'chiriladi)
               --rework            mavjud qoralamalarni (draft/needs_input) matn va kartochkasi bilan qayta yozish
               --rework=ID,ID      faqat shu postlarni
               --force-files       mashq fayllarini qayta yaratish
               --cards=PAPKA       faqat kartochkalarni shu papkaga chizish (post yozilmaydi) — ko'rib chiqish uchun

Kerak: Pillow (pip install pillow). Mashq fayllari uchun: openpyxl.
"""
import hashlib
import html
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
CH_KEY = "ai"
FONTS = Path(__file__).resolve().parent / "fonts"

DAYS = {"dush": 0, "sesh": 1, "chor": 2, "pay": 3, "juma": 4, "shan": 5, "yak": 6}
PANEL_KEYS = ("status", "message_id", "published_at", "published_by", "error", "publish_now",
              "attach_tries", "attach_error")


# ====================================================================== sozlamalar
def read_json(p: Path, default=None):
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return default


def channel() -> dict:
    """channels.json dagi «ai» kanali: {key, name, dir}."""
    for c in (read_json(ROOT / "channels.json", {}) or {}).get("channels", []):
        if c.get("key") == CH_KEY:
            return c
    raise SystemExit("channels.json da «ai» kanali topilmadi")


def channel_dir() -> Path:
    return ROOT / channel()["dir"]


def settings() -> dict:
    return read_json(channel_dir() / "settings.json", {}) or {}


# ====================================================================== matn yordamchilari
def esc(s) -> str:
    return html.escape(str(s), quote=False)


def pre(text: str, keep: bool = False) -> str:
    """Nusxa olinadigan so'rov bloki (Telegram <pre> — bosilsa nusxalanadi).
    Telefonda blok tor: qo'lda bo'lingan qatorlar chala-chulpa ko'rinadi. Shuning uchun qatorlar gap oxirida
    (. : ? !) bo'linadi, qolganlari birlashtiriladi — Telegram o'zi sig'diradi. keep=True — qatorlarga tegilmaydi."""
    lines = [l.strip() for l in text.strip("\n").split("\n")]
    if not keep:
        out = []
        for l in lines:
            if out and out[-1] and l and out[-1][-1] not in ".:?!" and not re.match(r"(?:[-•–]|\d+[.)])\s", l):
                out[-1] += " " + l
            else:
                out.append(l)
        lines = out
    return "<pre>" + esc("\n".join(lines)) + "</pre>"


def vislen(s: str) -> int:
    """Telegram sanaydigan uzunlik: teglarsiz, UTF-16 birliklarida."""
    t = html.unescape(re.sub(r"<[^>]+>", "", s))
    t = "\n".join(l for l in t.split("\n") if "{{APP_URL}}" not in l and "{{CONSULT_URL}}" not in l)
    return len(t.encode("utf-16-le")) // 2


def uzn(x, dec=0) -> str:
    """O'zbekcha son: 1 250 000; 12,5"""
    s = f"{x:,.{dec}f}".replace(",", " ").replace(".", ",")
    return s


# ====================================================================== kartochka chizgich
CW, CHH, SS = 1080, 1350, 2          # o'lcham va silliqlash koeffitsiyenti (2x chizib, kichraytiriladi)
PAPER, INK, INK2, INK3 = "#F5F2E9", "#141C2B", "#566070", "#8A8F98"
LINE, PANEL, DOT = "#D8D2C2", "#FFFFFF", "#CFC8B6"
RUBRICS = {
    # tag: (kartochkadagi yozuv, rang)
    "#sinov": ("HAFTA SINOVI", "#2F5BEA"),
    "#dars": ("DARS", "#0B8457"),
    "#ustalik": ("USTALIK", "#6D3FD0"),
    "#ehtiyot": ("EHTIYOT BO'LING", "#D7263D"),
    "#bonus": ("BONUS", "#D97706"),
}
DEFAULT_RUBRIC = ("AI DARSLAR", "#2F5BEA")
F_HEAD, F_TEXT, F_BOLD, F_MONO, F_MONOB = ("BricolageGrotesque-Bold.ttf", "InstrumentSans-Regular.ttf",
                                           "InstrumentSans-Bold.ttf", "JetBrainsMono-Regular.ttf",
                                           "JetBrainsMono-Bold.ttf")
_fonts, _cmaps = {}, {}


def _font(name: str, size: float) -> ImageFont.FreeTypeFont:
    key = (name, round(size * SS))
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype(str(FONTS / name), key[1])
    return _fonts[key]


def _has(name: str, text: str) -> bool:
    """Shriftda matnning hamma belgisi bormi (yo'q belgi «.notdef» shaklida chiqadi)."""
    if name not in _cmaps:
        f = ImageFont.truetype(str(FONTS / name), 24)
        nd = f.getmask("\uffff")
        _cmaps[name] = (f, (nd.size, bytes(nd)))
    f, notdef = _cmaps[name]
    for ch in set(text):
        if ch.isspace():
            continue
        m = f.getmask(ch)
        if (m.size, bytes(m)) == notdef:
            return False
    return True


def _clean(s) -> str:
    """Kartochka matni: shriftda yo'q tutuq belgilarini oddiy apostrofga almashtiradi."""
    return (str(s).replace("ʻ", "'").replace("’", "'").replace("‘", "'")
            .replace("ʼ", "'").replace("`", "'"))


class Card:
    """Bitta kartochka. Koordinatalar 1080×1350 o'lchovida beriladi."""

    M = 76  # chet bo'shliq

    def __init__(self, rubric: str, accent: str = None):
        label, color = RUBRICS.get(rubric, DEFAULT_RUBRIC)
        self.label, self.accent = label, accent or color
        self.img = Image.new("RGB", (CW * SS, CHH * SS), PAPER)
        self.d = ImageDraw.Draw(self.img)

    # ---------- past darajali
    @staticmethod
    def px(v):
        return int(round(v * SS))

    def font(self, name, size, text=""):
        if text and name not in (F_MONO, F_MONOB) and not _has(name, text):
            name = F_MONOB if name in (F_HEAD, F_BOLD) else F_MONO   # masalan kirill yozuv
        return _font(name, size)

    def tw(self, text, font, tracking=0.0):
        return self.d.textlength(text, font=font) / SS + tracking * max(0, len(text) - 1)

    def text(self, x, y, s, font, fill=INK, anchor="ls", tracking=0.0):
        """anchor: Pillow langari (ls — chap/asos chizig'i, rs, ms, lm, mm …)."""
        s = _clean(s)
        if not tracking:
            self.d.text((self.px(x), self.px(y)), s, font=font, fill=fill, anchor=anchor)
            return
        total = self.tw(s, font, tracking)
        cx = x - (total if anchor[0] == "r" else total / 2 if anchor[0] == "m" else 0)
        for ch in s:
            self.d.text((self.px(cx), self.px(y)), ch, font=font, fill=fill, anchor="l" + anchor[1])
            cx += self.d.textlength(ch, font=font) / SS + tracking

    def wrap(self, s, font, maxw):
        lines, cur = [], ""
        for word in _clean(s).split():
            t = (cur + " " + word).strip()
            if cur and self.tw(t, font) > maxw:
                lines.append(cur)
                cur = word
            else:
                cur = t
        if cur:
            lines.append(cur)
        return lines or [""]

    def fit(self, s, name, sizes, maxw, maxlines):
        """Eng katta o'lchamni tanlaydi: matn maxlines qatorga va maxw kenglikka sig'sin."""
        for size in sizes:
            f = self.font(name, size, s)
            lines = self.wrap(s, f, maxw)
            if len(lines) <= maxlines and all(self.tw(l, f) <= maxw for l in lines):
                return f, lines, size
        f = self.font(name, sizes[-1], s)
        lines = self.wrap(s, f, maxw)[:maxlines]
        while self.tw(lines[-1] + "…", f) > maxw and len(lines[-1]) > 1:
            lines[-1] = lines[-1][:-1]
        if len(self.wrap(s, f, maxw)) > maxlines:
            lines[-1] = lines[-1].rstrip() + "…"
        return f, lines, sizes[-1]

    def rect(self, x0, y0, x1, y1, fill=None, outline=None, r=0, width=1):
        box = [self.px(x0), self.px(y0), self.px(x1), self.px(y1)]
        if r:
            self.d.rounded_rectangle(box, radius=self.px(r), fill=fill, outline=outline, width=self.px(width))
        else:
            self.d.rectangle(box, fill=fill, outline=outline, width=self.px(width))

    def line(self, x0, y0, x1, y1, fill=LINE, width=2, dash=None):
        if not dash:
            self.d.line([self.px(x0), self.px(y0), self.px(x1), self.px(y1)], fill=fill, width=self.px(width))
            return
        on, off = dash
        length = ((x1 - x0) ** 2 + (y1 - y0) ** 2) ** 0.5
        if not length:
            return
        ux, uy, pos = (x1 - x0) / length, (y1 - y0) / length, 0.0
        while pos < length:
            end = min(pos + on, length)
            self.d.line([self.px(x0 + ux * pos), self.px(y0 + uy * pos), self.px(x0 + ux * end), self.px(y0 + uy * end)],
                        fill=fill, width=self.px(width))
            pos += on + off

    def circle(self, cx, cy, r, fill=None, outline=None, width=1):
        self.d.ellipse([self.px(cx - r), self.px(cy - r), self.px(cx + r), self.px(cy + r)],
                       fill=fill, outline=outline, width=self.px(width))

    # ---------- umumiy qolip: tepada rukn, sarlavha, pastda «Olib keting»
    def frame(self, title, meta="", sub=None, take=None, brand=""):
        M = self.M
        # rukn muhri
        f = self.font(F_MONOB, 25)
        w = self.tw(self.label, f, 2.5) + 44
        self.rect(M, 72, M + w, 72 + 54, fill=self.accent, r=9)
        self.text(M + 22, 72 + 27, self.label, f, fill="#FFFFFF", anchor="lm", tracking=2.5)
        if meta:
            self.text(CW - M, 72 + 27, meta, self.font(F_MONO, 24, meta), fill=INK2, anchor="rm", tracking=1.2)
        # sarlavha
        f, lines, size = self.fit(title, F_HEAD, [92, 84, 76, 68, 60, 54], CW - 2 * M, 3)
        y = 72 + 54 + 46 + size * 0.80
        for l in lines:
            self.text(M, y, l, f, fill=INK)
            y += size * 1.05
        y -= size * 1.05 - size * 0.24      # oxirgi qator tagi
        if sub:
            fs, sl, ss = self.fit(sub, F_TEXT, [36, 32, 29], CW - 2 * M, 2)
            y += 20 + ss * 0.8
            for l in sl:
                self.text(M, y, l, fs, fill=INK2)
                y += ss * 1.3
            y -= ss * 1.3 - ss * 0.25
        top = y + 44
        # past qism
        bottom = CHH - M
        if take:
            fv, vl, vs = self.fit(take, F_BOLD, [34, 31, 28, 25], CW - 2 * M, 2)
            vy = bottom - (len(vl) - 1) * vs * 1.28 - vs * 0.22
            for i, l in enumerate(vl):
                self.text(M, vy + i * vs * 1.28, l, fv, fill=INK)
            ly = vy - vs * 0.8 - 22
        else:
            ly = bottom - 6
        fl = self.font(F_MONOB, 21)
        if take:
            self.text(M, ly, "OLIB KETING", fl, fill=self.accent, tracking=2.2)
        if brand:
            self.text(CW - M, ly, brand, self.font(F_MONO, 23, brand), fill=INK2, anchor="rs")
        hy = ly - 21 * 0.8 - 26
        self.line(M, hy, CW - M, hy, fill=LINE, width=2)
        return top, hy - 40      # tana maydoni: yuqori va pastki chegarasi

    # ---------- tana turlari
    def body_stats(self, top, bot, stats, dots=None, caption=None):
        """Katta raqamlar + nuqtalar matritsasi. stats: [(qiymat, yozuv, accent?)], dots: {n, cols, marked:set}."""
        M = self.M
        n = len(stats)
        colw = (CW - 2 * M) / n
        size = 208 if n <= 2 else 150
        fb = self.font(F_HEAD, size)
        dots_h = 0
        if dots:
            cols = dots.get("cols", 50)
            rows = -(-dots["n"] // cols)
            cell = (CW - 2 * M) / cols
            dots_h = rows * cell
        cap_h = 40 if caption else 0
        block = size * 0.74 + 20 + 44 + (60 + dots_h if dots else 0) + (26 + cap_h if caption else 0)
        y = top + max(0, (bot - top - block) / 2) + size * 0.74
        for i, st in enumerate(stats):
            val, lab = str(st[0]), st[1]
            col = self.accent if (len(st) > 2 and st[2]) else INK
            x = M + i * colw
            self.text(x - size * 0.03, y, val, fb, fill=col)
            fl, ll, ls = self.fit(lab, F_TEXT, [40, 36, 32, 28], colw - 30, 2)
            for j, l in enumerate(ll):
                self.text(x, y + 20 + ls * 0.8 + j * ls * 1.25, l, fl, fill=INK2)
        y += 20 + 44 + 30
        if dots:
            y += 30
            marked = dots.get("marked") or set()
            r = cell * 0.31
            for k in range(dots["n"]):
                cx = M + (k % cols) * cell + cell / 2
                cy = y + (k // cols) * cell + cell / 2
                self.circle(cx, cy, r, fill=self.accent if k in marked else DOT)
            y += dots_h + 26
        if caption:
            self.text(M, y + 24, caption, self.font(F_MONO, 25, caption), fill=INK3, tracking=0.6)

    def body_chain(self, top, bot, steps):
        """Qadamlar zanjiri. steps: [(sarlavha, izoh)] — izoh mono shriftda (so'rovning kalit iborasi)."""
        M, n = self.M, len(steps)
        rowh = min((bot - top) / n, 176)
        y0 = top + ((bot - top) - rowh * n) / 2
        r, cx = 34, M + 34
        self.line(cx, y0 + rowh / 2, cx, y0 + rowh * (n - 0.5), fill=self.accent, width=4)
        tx, maxw = M + 2 * r + 30, CW - 2 * M - 2 * r - 30
        for i, st in enumerate(steps):
            title, note = (st if isinstance(st, (tuple, list)) else (st, None))
            cy = y0 + rowh * (i + 0.5)
            last = i == n - 1
            self.circle(cx, cy, r, fill=self.accent if last else PAPER, outline=self.accent, width=4)
            self.text(cx, cy + 1, str(i + 1), self.font(F_MONOB, 31), fill="#FFFFFF" if last else self.accent, anchor="mm")
            ft, tl, ts = self.fit(title, F_BOLD, [46, 42, 38, 34, 30], maxw, 1)
            if note:
                fn, nl, ns = self.fit(note, F_MONO, [28, 26, 24, 22, 20], maxw, 1)
                self.text(tx, cy - 7, tl[0], ft, fill=INK)
                self.text(tx, cy + 12 + ns * 0.85, nl[0], fn, fill=INK2)
            else:
                self.text(tx, cy, tl[0], ft, fill=INK, anchor="lm")

    def body_pairs(self, top, bot, left, right, links, heads=("", ""), legend=True):
        """Ikki ro'yxatni solishtirish. links: [(chap_i, o'ng_j, "=" aniq | "~" o'xshash)]."""
        M = self.M
        n = max(len(left), len(right))
        leg_h = 76 if legend else 0
        head_h = 46
        rowh = min((bot - top - leg_h - head_h) / n, 106)
        y0 = top + ((bot - top - leg_h - head_h) - rowh * n) / 2 + head_h
        colw = 392
        xl, xr = M, CW - M - colw
        fh = self.font(F_MONOB, 23)
        self.text(xl, y0 - 22, heads[0], fh, fill=INK2, tracking=2.2)
        self.text(xr, y0 - 22, heads[1], fh, fill=INK2, tracking=2.2)
        cellh = rowh - 18
        linked_l = {a: k for a, b, k in links}
        linked_r = {b: k for a, b, k in links}

        def cell(x, i, s, state):
            y = y0 + i * rowh
            col = {"=": self.accent, "~": self.accent, None: LINE}[state]
            self.rect(x, y, x + colw, y + cellh, fill=PANEL, outline=col, r=10, width=2 if state else 2)
            f, ls, _ = self.fit(s, F_MONO, [27, 25, 23, 21, 19], colw - 36, 1)
            self.text(x + 18, y + cellh / 2, ls[0], f, fill=INK if state else INK3, anchor="lm")
            return y + cellh / 2

        ly = [cell(xl, i, s, linked_l.get(i)) for i, s in enumerate(left)]
        ry = [cell(xr, j, s, linked_r.get(j)) for j, s in enumerate(right)]
        for a, b, k in links:
            self.line(xl + colw, ly[a], xr, ry[b], fill=self.accent, width=3, dash=None if k == "=" else (9, 8))
            self.circle(xl + colw, ly[a], 7, fill=self.accent)
            self.circle(xr, ry[b], 7, fill=self.accent)
        if legend:
            y = bot - 10
            f = self.font(F_TEXT, 27)
            x = M
            for dash, lab in ((None, "aniq mos"), ((9, 8), "o'xshash — siz hal qilasiz"), ("box", "faqat bittasida")):
                if dash == "box":
                    self.rect(x, y - 16, x + 44, y + 4, fill=PANEL, outline=LINE, r=5, width=2)
                else:
                    self.line(x, y - 7, x + 44, y - 7, fill=self.accent, width=3, dash=dash)
                self.text(x + 56, y, lab, f, fill=INK2)
                x += 56 + self.tw(lab, f) + 30

    def body_checks(self, top, bot, items):
        """Tekshiruv ro'yxati. items: [(tuzoq, nazorat savoli)]."""
        M, n = self.M, len(items)
        rowh = min((bot - top) / n, 176)
        y0 = top + ((bot - top) - rowh * n) / 2
        for i, (trap, ask) in enumerate(items):
            y = y0 + i * rowh
            if i:
                self.line(M, y, CW - M, y, fill=LINE, width=2)
            cy = y + rowh / 2
            self.text(M, cy - 7, f"{i + 1:02d}", self.font(F_MONOB, 40), fill=self.accent)
            tx, maxw = M + 96, CW - 2 * M - 96
            ft, tl, ts = self.fit(trap, F_BOLD, [42, 39, 36, 33, 30], maxw, 1)
            self.text(tx, cy - 7, tl[0], ft, fill=INK)
            fa, al, as_ = self.fit("→ " + ask, F_MONO, [27, 25, 23, 21], maxw, 1)
            self.text(tx, cy + 13 + as_ * 0.85, al[0], fa, fill=INK2)

    def body_bars(self, top, bot, items, note=None, unit=""):
        """Gorizontal ustunlar. items: [(yozuv, qiymat, ajratilganmi)], note: (katta matn, izoh)."""
        M, n = self.M, len(items)
        note_h = 190 if note else 0
        rowh = min((bot - top - note_h) / n, 104)
        y0 = top + ((bot - top - note_h) - rowh * n) / 2
        mx = max(v for _, v, *_ in items) or 1
        labw = 330
        fv = self.font(F_MONOB, 28)
        valw = max(self.tw(uzn(v) + unit, fv) for _, v, *_ in items) + 18
        barw = CW - 2 * M - labw - valw
        for i, it in enumerate(items):
            lab, v, hot = it[0], it[1], (len(it) > 2 and it[2])
            cy = y0 + rowh * (i + 0.5)
            fl, ll, _ = self.fit(lab, F_BOLD if hot else F_TEXT, [34, 31, 28, 25], labw - 16, 1)
            self.text(M, cy, ll[0], fl, fill=INK if hot else INK2, anchor="lm")
            w = max(10, barw * v / mx)
            h = min(rowh * 0.5, 46)
            self.rect(M + labw, cy - h / 2, M + labw + w, cy + h / 2, fill=self.accent if hot else DOT, r=6)
            self.text(M + labw + w + 14, cy, uzn(v) + unit, fv, fill=INK if hot else INK2, anchor="lm")
        if note:
            big, small = note
            y = bot - note_h + 36
            self.rect(M, y, CW - M, bot, fill=PANEL, outline=LINE, r=14, width=2)
            mid = y + (bot - y) / 2
            fb, bl, bs = self.fit(big, F_HEAD, [60, 54, 48, 42, 36], CW - 2 * M - 68, 1)
            self.text(M + 34, mid + 2, bl[0], fb, fill=self.accent)
            fs, sl, ss = self.fit(small, F_TEXT, [28, 26, 24, 22], CW - 2 * M - 68, 1)
            self.text(M + 34, mid + 22 + ss * 0.85, sl[0], fs, fill=INK2)

    def body_bullets(self, top, bot, items):
        """Oddiy ro'yxat (umumiy muqova uchun). items: [satr]."""
        M, n = self.M, len(items)
        rowh = min((bot - top) / n, 140)
        y0 = top + ((bot - top) - rowh * n) / 2
        for i, s in enumerate(items):
            cy = y0 + rowh * (i + 0.5)
            self.rect(M, cy - 9, M + 18, cy + 9, fill=self.accent, r=4)
            f, ls, size = self.fit(s, F_TEXT, [42, 38, 34, 30], CW - 2 * M - 50, 2)
            yy = cy - (len(ls) - 1) * size * 0.62
            for j, l in enumerate(ls):
                self.text(M + 46, yy + j * size * 1.24, l, f, fill=INK, anchor="lm")

    def body_prompt(self, top, bot, lines, label="SO'ROV"):
        """Terminal uslubidagi so'rov bloki. lines: [satr]."""
        M = self.M
        size = 30
        f = self.font(F_MONO, size)
        maxw = CW - 2 * M - 72
        out = []
        for l in lines:
            out += self.wrap(l, f, maxw) if l.strip() else [""]
        lh = size * 1.5
        h = min(bot - top, len(out) * lh + 120)
        y = top + ((bot - top) - h) / 2
        self.rect(M, y, CW - M, y + h, fill=INK, r=18)
        self.text(M + 36, y + 50, label, self.font(F_MONOB, 21), fill=self.accent if self.accent != INK else "#FFFFFF", tracking=2.2)
        yy = y + 96
        for l in out:
            if yy > y + h - 24:
                break
            self.text(M + 36, yy, l, f, fill="#E9EDF3")
            yy += lh

    def png(self) -> Image.Image:
        return self.img.resize((CW, CHH), Image.LANCZOS)


def render_card(spec: dict, rubric: str, title: str, meta: str, brand: str, take: str = None) -> Image.Image:
    """spec: {"kind": "stats"|"chain"|"pairs"|"checks"|"bars"|"bullets"|"prompt", "title"?, "sub"?, "take"?, …}"""
    c = Card(rubric, spec.get("accent"))
    top, bot = c.frame(spec.get("title") or title, meta=meta, sub=spec.get("sub"),
                       take=spec.get("take", take), brand=brand)
    kind = spec["kind"]
    if kind == "stats":
        c.body_stats(top, bot, spec["stats"], spec.get("dots"), spec.get("caption"))
    elif kind == "chain":
        c.body_chain(top, bot, spec["steps"])
    elif kind == "pairs":
        c.body_pairs(top, bot, spec["left"], spec["right"], spec["links"], spec.get("heads", ("", "")))
    elif kind == "checks":
        c.body_checks(top, bot, spec["items"])
    elif kind == "bars":
        c.body_bars(top, bot, spec["items"], spec.get("note"), spec.get("unit", ""))
    elif kind == "bullets":
        c.body_bullets(top, bot, spec["items"])
    elif kind == "prompt":
        c.body_prompt(top, bot, spec["lines"], spec.get("label", "SO'ROV"))
    else:
        raise SystemExit(f"Noma'lum kartochka turi: {kind}")
    return c.png()


def gpt_prompt(rubric: str, headline: str, scene: str, brand: str) -> str:
    """Ixtiyoriy: kartochka o'rniga AI rasm xohlansa — GPT Image uchun prompt (kanal uslubida)."""
    label, color = RUBRICS.get(rubric, DEFAULT_RUBRIC)
    return f"""Create a vertical 4:5 cover image for a Telegram channel post. Style: clean paper-cut editorial illustration — flat layered shapes with soft shadows, like a neatly filled-in office form. Warm off-white paper background ({PAPER}), deep navy ink ({INK}) and one accent colour ({color}). Calm, trustworthy, no clutter.

1) Top-left: a small solid {color} rounded rectangle with white bold uppercase monospace text "{label}".
2) Center: {scene} Paper-cut style, matte, generous empty space around it. No people's faces, no real logos, no state emblems or flags.
3) Below the illustration: a large bold navy headline in a characterful grotesque sans-serif, at most three lines: "{_clean(headline)}"
4) Bottom, above a thin hairline: left small monospace navy text "{brand}".

Render every piece of text EXACTLY as written, spelled correctly, with straight apostrophes. No other text, no watermarks."""


# ====================================================================== hafta yig'uvchi
class Week:
    def __init__(self, n: int, start: str, theme: str):
        self.n, self.theme = n, theme
        self.args = sys.argv[1:]
        self.moved = False
        for a in self.args:
            if a.startswith("--start="):
                start, self.moved = a.split("=", 1)[1], True
        self.start = date.fromisoformat(start)
        if self.start.weekday() != 0:
            raise SystemExit(f"Hafta dushanbadan boshlanadi — {start} dushanba emas")
        self.ch = channel()
        self.base = ROOT / self.ch["dir"]
        self.cfg = settings()
        handle = str(self.cfg.get("channel_id") or "").strip()
        self.brand = handle if handle.startswith("@") else (self.ch.get("name") or "AI darslar")
        self.time = str(self.cfg.get("post_time") or "08:30")
        self.files_rel = f"files/h{n:02d}"
        self.posts = []

    # ---------- sana va ID
    def day(self, key: str) -> date:
        return self.start + timedelta(days=DAYS[key])

    def pid(self, key: str, time: str = None) -> str:
        t = (time or self.time).replace(":", "")
        return f"{self.day(key).isoformat()}-{t}"

    # ---------- mashq fayllari
    def file(self, name: str, builder) -> str:
        """Mashq faylini channels/ai/files/hNN/ ga yaratadi (bor bo'lsa tegmaydi) va nisbiy yo'lini qaytaradi.
        builder(path) — faylni yozadigan funksiya. Fayl nomi: lotin harflari, raqam, _ - ."""
        if not re.fullmatch(r"[A-Za-z0-9_\-.]{1,80}", name):
            raise SystemExit(f"Fayl nomi noto'g'ri: {name}")
        rel = f"{self.files_rel}/{name}"
        path = self.base / rel
        if path.exists() and "--force-files" not in self.args:
            return rel
        path.parent.mkdir(parents=True, exist_ok=True)
        builder(path)
        print(f"fayl: {rel}  ({path.stat().st_size // 1024} KB)")
        return rel

    # ---------- post
    def post(self, day: str, rubric: str, title: str, head: str, body: str, new: str, take: str, card: dict,
             emoji: str = "", attach: list = None, notes: str = "", sources: list = None, scene: str = None,
             time: str = None):
        """day — dush|sesh|chor|pay|juma|shan|yak. new — «Bugungi yangi bilim», take — «Olib keting» (ikkalasi majburiy).
        card — kartochka tavsifi (render_card). scene — ixtiyoriy AI rasm prompti uchun sahna tavsifi (inglizcha)."""
        if not (new or "").strip() or not (take or "").strip():
            raise SystemExit(f"{day}: har postda «Bugungi yangi bilim» (new) va «Olib keting» (take) bo'lishi shart")
        pid = self.pid(day, time)
        parts = [f"{emoji} <b>{head}</b>".strip(), "", body.strip(), "",
                 f"💡 <b>Bugungi yangi bilim:</b> {new.strip()}", "",
                 f"🎁 <b>Olib keting:</b> {take.strip()}", "",
                 "➖➖➖", "{{CHANNEL}}  " + rubric]
        caption = "\n".join(parts)
        n = vislen(caption)
        if n > 4096:
            raise SystemExit(f"{pid}: matn {n} belgi — Telegram chegarasi 4096. Qisqartiring yoki ikki postga bo'ling.")
        p = dict(id=pid, rubric=rubric, title=title, scheduled_at=f"{pid[:10]}T{pid[11:13]}:{pid[13:15]}:00+05:00",
                 caption=caption, prompt=gpt_prompt(rubric, card.get("title") or head, scene, self.brand) if scene else "",
                 product_url=None, notes=notes, sources=sources or [], status="draft",
                 week=self.n, attachments=attach or [])
        if n > 1024:
            p["long_ok"] = True      # dars uzun: rasm va matn ikki xabar bo'lib chiqadi — ataylab
        self.posts.append((p, card, head, take))

    # ---------- yozish
    def build(self):
        posts_dir, img_dir = self.base / "posts", self.base / "images"
        posts_dir.mkdir(parents=True, exist_ok=True)
        img_dir.mkdir(parents=True, exist_ok=True)
        self.posts.sort(key=lambda x: x[0]["scheduled_at"])
        total = len(self.posts)
        cards_only = next((a.split("=", 1)[1] for a in self.args if a.startswith("--cards=")), None)
        rework = None
        for a in self.args:
            if a == "--rework":
                rework = set()
            elif a.startswith("--rework="):
                rework = {x for x in a.split("=", 1)[1].split(",") if x}
        ids = {p["id"] for p, *_ in self.posts}

        if self.moved and not cards_only:     # boshqa haftaga ko'chirildi — eski qoralamalarni tozalash
            for f in sorted(posts_dir.glob("*.json")):
                cur = read_json(f, {}) or {}
                if cur.get("week") != self.n or cur.get("id") in ids:
                    continue
                if cur.get("status") in ("draft", "needs_input"):
                    if cur.get("image_auto") and cur.get("image"):
                        (img_dir / cur["image"]).unlink(missing_ok=True)
                    f.unlink()
                    print(f"{cur.get('id')}  eski qoralama o'chirildi (hafta ko'chirildi)")
                else:
                    print(f"{cur.get('id')}  {cur.get('status')} — o'chirilmadi, panelda o'zingiz hal qiling")

        print(f"\n{self.n}-hafta · {self.theme} · {self.start:%d.%m.%Y} dan · kartochkada: {self.brand}\n")
        for i, (p, card, head, take) in enumerate(self.posts, 1):
            meta = f"{self.n}-HAFTA · {i}/{total}"
            img = render_card(card, p["rubric"], head, meta, self.brand, take=card.get("take", _short_take(take)))
            if cards_only:
                out = Path(cards_only)
                out.mkdir(parents=True, exist_ok=True)
                img.save(out / f"{p['id']}.png", optimize=True)
                print(f"{p['id']}  kartochka -> {out / (p['id'] + '.png')}")
                continue
            target = posts_dir / f"{p['id']}.json"
            n = vislen(p["caption"])
            if target.exists():
                cur = read_json(target, {}) or {}
                if rework is None or (rework and p["id"] not in rework):
                    print(f"{p['id']}  mavjud — tegilmadi (qayta yozish: --rework)")
                    continue
                if cur.get("status") not in ("draft", "needs_input"):
                    print(f"{p['id']}  {cur.get('status')} — qayta yozilmadi")
                    continue
                for k in PANEL_KEYS:                       # panel maydonlari saqlanadi
                    if k in cur:
                        p[k] = cur[k]
                if cur.get("image") and not cur.get("image_auto"):
                    p["image"] = cur["image"]              # qo'lda yuklangan rasmga tegilmaydi
                sent = {a.get("file"): a.get("message_id") for a in cur.get("attachments") or [] if a.get("message_id")}
                for a in p["attachments"]:
                    if sent.get(a["file"]):
                        a["message_id"] = sent[a["file"]]
            if "image" not in p:
                name = f"{p['id']}-auto.png"
                img.save(img_dir / name, optimize=True)
                p["image"], p["image_auto"] = name, True
            target.write_text(json.dumps(p, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            files = ", ".join(a["file"].split("/")[-1] for a in p["attachments"]) or "—"
            print(f"{p['id']}  {p['rubric']:9} {n:5}  {p['title']}  | fayl: {files}")
        print()


def _short_take(take: str) -> str:
    """«Olib keting» matnining kartochkaga sig'adigan qismi (teglarsiz, birinchi gap)."""
    t = html.unescape(re.sub(r"<[^>]+>", "", take)).strip()
    t = re.split(r"(?<=[.!?])\s", t)[0].rstrip(".")
    return t


# ====================================================================== Excel yordamchisi
def write_table(ws, headers, rows, widths=None, title=None, note=None, date_cols=(), money_cols=(), text_cols=()):
    """Oddiy jadval: (ixtiyoriy) birlashtirilgan sarlavha qatori, qalin ustun nomlari, kengliklar.
    Qaytaradi: ustun nomlari turgan qator raqami."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    thin = Side(style="thin", color="BFBFBF")
    r = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(name="Arial", size=12, bold=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
        r = 2
        if note:
            ws.cell(row=2, column=1, value=note).font = Font(name="Arial", size=9, italic=True, color="7F7F7F")
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
            ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")
            r = 3
    head_row = r
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font, cell.fill = bold, PatternFill("solid", fgColor="E7E6E6")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = font
            if c in date_cols and not isinstance(v, str) and v is not None:
                cell.number_format = "DD.MM.YYYY"
            if c in money_cols and isinstance(v, (int, float)):
                cell.number_format = "#,##0"
            if c in text_cols:
                cell.number_format = "@"
    for c, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    return head_row


def stable_seed(*parts) -> int:
    """Tasodifiy ma'lumot har safar bir xil chiqishi uchun (random.Random(stable_seed(...)))."""
    return int(hashlib.sha256("|".join(map(str, parts)).encode()).hexdigest()[:12], 16)
