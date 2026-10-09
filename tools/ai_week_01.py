# -*- coding: utf-8 -*-
"""«AI darslar» — 1-hafta: Excel faylni kod bilan tahlil qildirish.

    python tools/ai_week_01.py                      # postlar, kartochkalar, mashq fayllari
    python tools/ai_week_01.py --start=2026-10-19   # boshqa haftaga ko'chirish
    python tools/ai_week_01.py --rework             # qoralamalarni qayta yozish
    python tools/ai_week_01.py --cards=/tmp/k       # faqat kartochkalarni ko'rish

Bu fayl keyingi haftalar uchun NAMUNA: nusxa oling (ai_week_02.py), mavzuni, mashq fayllarini va 5 ta
postni almashtiring. Postdagi har bir raqam shu yerda ma'lumotdan hisoblanadi — qo'lda yozilmaydi.
"""
import random
from collections import Counter
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ai_lib import Week, pre, stable_seed, uzn, write_table

W = Week(1, start="2026-10-12", theme="Excel faylni kod bilan tahlil qildirish")
CHECKED = "09.10.2026"      # vositalar imkoniyati tekshirilgan sana (har hafta yangilanadi)

SRC_TOOLS = [
    {"label": f"ChatGPT: fayl yuklash va cheklovlar — bepul tarifda kuniga 3 ta fayl ({CHECKED} holatiga)",
     "url": "https://help.openai.com/en/articles/8555545-file-uploads-faq-advanced-data-analysis"},
    {"label": f"Claude: kod bajarish va fayl yaratish — hamma tarifda, bepulida ham ({CHECKED} holatiga)",
     "url": "https://support.claude.com/en/articles/12111783-create-and-edit-files-with-claude"},
    {"label": f"Gemini: fayl yuklash va tahlil ({CHECKED} holatiga)",
     "url": "https://support.google.com/gemini/answer/14903178?hl=en"},
]

OYLAR = ["yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul", "avgust", "sentabr", "oktabr", "noyabr", "dekabr"]


# ══════════════════════════════════════════════════════════════════ 1-fayl: murojaatlar reestri
REGIONS = [("Toshkent shahri", 20), ("Toshkent viloyati", 9), ("Samarqand viloyati", 10), ("Farg'ona viloyati", 8),
           ("Andijon viloyati", 7), ("Namangan viloyati", 7), ("Qashqadaryo viloyati", 6), ("Buxoro viloyati", 6),
           ("Surxondaryo viloyati", 5), ("Xorazm viloyati", 5), ("Jizzax viloyati", 4), ("Navoiy viloyati", 4),
           ("Sirdaryo viloyati", 3), ("Qoraqalpog'iston Respublikasi", 6)]
DIRTY = {"Toshkent shahri": ["Toshkent sh.", "г.Ташкент", "toshkent shahri", "Toshkent shahri "],
         "Samarqand viloyati": ["Samarqand vil.", "Самаркандская обл."]}
# (mavzu, mas'ul bo'lim, muddat kun, ulush)
TOPICS = [("Ruxsatnoma olish tartibi", "Ruxsatnomalar bo'limi", 15, 12),
          ("Ruxsatnoma muddatini uzaytirish", "Ruxsatnomalar bo'limi", 30, 9),
          ("To'lov va qarzdorlik", "Moliya bo'limi", 15, 16),
          ("Ma'lumotnoma so'rash", "Fuqarolar bilan ishlash bo'limi", 15, 22),
          ("Taklif va tashakkur", "Fuqarolar bilan ishlash bo'limi", 15, 5),
          ("Qonunchilik bo'yicha tushuntirish", "Yuridik bo'lim", 15, 14),
          ("Xodim xatti-harakati ustidan shikoyat", "Nazorat bo'limi", 30, 12),
          ("Elektron xizmat ishlamayapti", "Axborot texnologiyalari bo'limi", 15, 6)]
SPIKE_TOPIC, SPIKE_MONTH, SPIKE_EXTRA = "Elektron xizmat ishlamayapti", 9, 40
WORST_DEPT = "Ruxsatnomalar bo'limi"
LATE_SHARE = {WORST_DEPT: 0.34}      # qolgan bo'limlarda — 0.07
MONTH_BASE = [38, 35, 39, 37, 36, 38, 37, 36, 37, 39, 37, 39]
AS_OF = date(2025, 12, 31)


def workdays(y, m):
    d, out = date(y, m, 1), []
    while d.month == m:
        if d.weekday() < 5:
            out.append(d)
        d += timedelta(days=1)
    return out


def next_workday(d):
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def gen_appeals():
    rnd = random.Random(stable_seed("h01", "murojaatlar"))
    rows = []
    for m, base in enumerate(MONTH_BASE, 1):
        days = workdays(2025, m)
        for _ in range(base):
            t = rnd.choices(TOPICS, weights=[x[3] for x in TOPICS])[0]
            rows.append({"date": rnd.choice(days), "topic": t[0], "dept": t[1], "term": t[2]})
        if m == SPIKE_MONTH:      # yashirin muammo №2: bir oyda bitta mavzu bo'yicha keskin o'sish
            t = next(x for x in TOPICS if x[0] == SPIKE_TOPIC)
            burst = [d for d in days if 8 <= d.day <= 26]
            for _ in range(SPIKE_EXTRA):
                rows.append({"date": rnd.choice(burst), "topic": t[0], "dept": t[1], "term": t[2], "spike": True})
    rows.sort(key=lambda r: (r["date"], rnd.random()))
    for r in rows:
        r["region"] = rnd.choices([x[0] for x in REGIONS], weights=[x[1] for x in REGIONS])[0]
        r["who"] = ("Y-" if rnd.random() < 0.3 else "F-") + str(rnd.randint(10000, 99999))
        r["open"] = r["date"] >= date(2025, 12, 18) and rnd.random() < 0.85
    # yashirin muammo №1: bitta bo'lim muddatni ko'p buzadi (aniq son — tasodifga bog'liq emas)
    for dept in sorted({r["dept"] for r in rows}):
        pool = [r for r in rows if r["dept"] == dept and not r["open"] and r["date"] < date(2025, 11, 20)]
        done = [r for r in rows if r["dept"] == dept and not r["open"]]
        for r in rnd.sample(pool, round(len(done) * LATE_SHARE.get(dept, 0.07))):
            r["late"] = True
    for r in rows:
        if r["open"]:
            r["answer"] = None
            continue
        for _ in range(200):
            days = rnd.randint(r["term"] + 1, r["term"] + 18) if r.get("late") else rnd.randint(2, r["term"])
            ans = r["date"] + timedelta(days=days)
            if ans.weekday() < 5 and ans <= AS_OF:
                break
        else:
            ans = min(next_workday(r["date"] + timedelta(days=3)), AS_OF)
        r["answer"] = ans
    # yashirin muammo №3: ayrim murojaatlar ikki marta ro'yxatga olingan
    originals = rnd.sample([i for i, r in enumerate(rows) if not r["open"] and not r.get("spike")], 12)
    for i in sorted(originals, reverse=True):
        copy = dict(rows[i], dup=True)
        rows[i]["dup_of"] = True
        rows.insert(min(len(rows), i + rnd.randint(1, 5)), copy)
    # «iflos» ma'lumot: hudud turlicha yozilgan, bo'sh kataklar, matn bo'lib qolgan sana
    for r in rows:
        r["region_raw"] = r["region"]
        if r["region"] in DIRTY and rnd.random() < 0.28:
            r["region_raw"] = rnd.choice(DIRTY[r["region"]])
    blanks = rnd.sample([r for r in rows if r["dept"] != WORST_DEPT and not r.get("dup") and not r.get("dup_of")], 7)
    for r in blanks:
        r["dept_raw"] = None
    amb = [r for r in rows if r["date"].day <= 12 and r["date"].day != r["date"].month and not r.get("dup") and not r.get("dup_of")]
    for r in rnd.sample(amb, 6):
        r["date_text"] = True
    for i, r in enumerate(rows, 1):
        r["reg"] = f"M-{i:04d}"
    return rows


APPEALS = gen_appeals()


def appeal_stats(rows):
    """Javoblar — tozalangan ma'lumot bo'yicha (takrorlar olib tashlangan), o'quvchi 2–3-qadamda oladigan natija."""
    uniq = [r for r in rows if not r.get("dup")]
    done = [r for r in uniq if r["answer"]]
    ontime = lambda rs: sum(1 for r in rs if (r["answer"] - r["date"]).days <= r["term"])
    by = {}
    for r in done:
        by.setdefault(r.get("dept_raw", r["dept"]) or "(bo'lim ko'rsatilmagan)", []).append(r)
    pct = {d: 100 * ontime(rs) / len(rs) for d, rs in by.items()}
    worst = min((d for d in pct if not d.startswith("(")), key=lambda d: pct[d])
    others = [r for r in done if r.get("dept_raw", r["dept"]) != worst]
    months = Counter(r["date"].month for r in uniq)
    top_m = max(months, key=months.get)
    rest = [v for k, v in months.items() if k != top_m]
    topic_m = Counter(r["topic"] for r in uniq if r["date"].month == top_m)
    late_w, late_o = 100 - pct[worst], 100 - 100 * ontime(others) / len(others)
    return dict(total=len(rows), uniq=len(uniq), dups=len(rows) - len(uniq), done=len(done), open=len(uniq) - len(done),
                overall=100 * ontime(done) / len(done), worst=worst, worst_pct=pct[worst], worst_n=len(by[worst]),
                others_pct=100 * ontime(others) / len(others), ratio=late_w / late_o, late_w=late_w, late_o=late_o,
                month=top_m, month_n=months[top_m], month_avg=sum(rest) / len(rest),
                month_topic=topic_m.most_common(1)[0], blanks=sum(1 for r in uniq if r.get("dept_raw", 1) is None),
                text_dates=sum(1 for r in rows if r.get("date_text")),
                dirty_regions=sum(1 for r in rows if r["region_raw"] != r["region"]))


A = appeal_stats(APPEALS)
assert A["total"] == 500 and A["worst"] == WORST_DEPT and A["month"] == SPIKE_MONTH and A["dups"] == 12, A
assert A["ratio"] >= 3 and A["overall"] >= 80, A      # «umumiy foizda ko'rinmaydi» sharti


def build_appeals(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reestr"
    rows = []
    for r in APPEALS:
        d = r["date"].strftime("%d.%m.%Y") if r.get("date_text") else r["date"]
        rows.append([r["reg"], d, r["who"], r["region_raw"], r["topic"], r.get("dept_raw", r["dept"]), r["term"],
                     r["answer"], "Bajarildi" if r["answer"] else "Ko'rib chiqilmoqda"])
    write_table(ws, ["Reg. №", "Kelib tushgan sana", "Murojaatchi kodi", "Hudud", "Murojaat mavzusi", "Mas'ul bo'lim",
                     "Muddat (kun)", "Javob berilgan sana", "Holat"], rows,
                widths=[10, 14, 14, 30, 38, 34, 10, 14, 20],
                title="«Namuna» agentligi: jismoniy va yuridik shaxslar murojaatlari reestri, 2025-yil",
                note="O'quv mashqi uchun o'ylab topilgan ma'lumot (31.12.2025 holatiga). Haqiqiy idora yoki shaxslarga aloqasi yo'q.",
                date_cols=(2, 8))
    wb.save(path)


F_APPEALS = W.file("murojaatlar_2025.xlsx", build_appeals)


# ══════════════════════════════════════════════════════════════════ 2–3-fayl: ombor va buxgalteriya
CATALOG = {
    "Noutbuk": ["HP 250 G8", "HP 250 G9", "HP ProBook 450 G9", "Lenovo IdeaPad 3", "Lenovo ThinkBook 15", "Acer Aspire 5", "Asus VivoBook 15", "Dell Vostro 3510"],
    "Monitor": ["LG 24MK430", "LG 27MP400", "Samsung S24R350", "Samsung S27A330", "Philips 243V7", "Acer K242HYL", "Dell SE2422H"],
    "Printer": ["Canon LBP6030", "Canon LBP2900", "HP LaserJet Pro M404dn", "HP LaserJet M111a", "Epson L3210", "Brother HL-L2340", "Pantum P2207"],
    "MFU": ["Canon MF3010", "HP LaserJet MFP M141a", "Canon MF445dw", "Epson L3250", "Kyocera M2040dn"],
    "Skaner": ["Canon LiDE 300", "Epson DS-530", "HP ScanJet Pro 2000"],
    "Proyektor": ["Epson EB-X51", "Epson EB-E01", "BenQ MS550", "ViewSonic PA503S"],
    "Kompyuter": ["Dell OptiPlex 3080", "HP ProDesk 400 G7", "Lenovo ThinkCentre M70s", "Acer Veriton X2690"],
    "Konditsioner": ["Artel ART-12HG", "Artel ART-18HM", "Samsung AR12", "Gree GWH12", "Midea MSAG-12"],
    "UPS": ["APC Back-UPS 650VA", "APC Back-UPS 1100VA", "Ippon Back Basic 850", "CyberPower UT850"],
    "Router": ["TP-Link Archer C6", "TP-Link TL-WR841N", "MikroTik hAP ac2", "Keenetic Giga"],
    "Televizor": ["Artel 43 dyuym", "Samsung 50 dyuym", "LG 55 dyuym"],
    "Telefon apparati": ["Panasonic KX-TS2350", "Panasonic KX-TG1611", "Yealink T31P"],
    "Stol": ["yozuv 140 sm", "yozuv 160 sm", "rahbar uchun 180 sm", "majlis 240 sm", "kompyuter uchun 120 sm"],
    "Kreslo": ["ofis, to'rli", "ofis, charm", "rahbar uchun", "mehmon uchun"],
    "Stul": ["majlis zali uchun", "yumshoq, qora", "plastik"],
    "Shkaf": ["hujjatlar uchun 2 eshikli", "hujjatlar uchun oynali", "kiyim uchun", "metall 4 tortmali"],
    "Seyf": ["60 sm", "90 sm", "120 sm, o'tga chidamli"],
    "Javon": ["kitoblar uchun", "arxiv uchun metall"],
    "Tumba": ["3 tortmali", "g'ildirakli"],
    "Doska": ["magnit-marker 120x90", "probkali 90x60", "flipchart"],
    "Kuler": ["HotFrost V118", "Artel suv dispenseri"],
    "Isitgich": ["moyli 9 seksiya", "konvektor 2 kVt"],
    "Shreder": ["Fellowes 60Cs", "Rexel Alpha", "Kobra +1 SS7"],
    "Laminator": ["Fellowes Saturn A3", "Office Kit L3225"],
    "Kalkulyator": ["Citizen SDC-888", "Casio GR-12"],
    "Veb-kamera": ["Logitech C270", "Logitech C920"],
    "Kolonka": ["Sven SPS-702", "Logitech Z333"],
    "Fotoapparat": ["Canon EOS 2000D"],
    "Muzlatgich": ["Artel HS 117", "Samsung RT22"],
    "Mikroto'lqinli pech": ["Artel 20 l", "Samsung ME83"],
    "Choynak": ["elektr, 1,7 l"],
    "Soat": ["devor, dumaloq"],
    "Gilam": ["3x4 m"],
    "Parda": ["jalyuzi, vertikal"],
    "O't o'chirgich": ["OP-5", "OU-3"],
    "Planshet": ["Samsung Galaxy Tab A8", "Lenovo Tab M10"],
    "Server": ["HP ProLiant ML30"],
    "Kommutator": ["D-Link DES-1016D", "TP-Link TL-SG1024D"],
    "Videokamera": ["Hikvision DS-2CD1023"],
    "Mikrofon": ["Shure SM58"],
    "Klaviatura": ["Logitech K120"],
    "Sichqoncha": ["Logitech M90"],
    "Tashqi disk": ["WD Elements 1TB", "Seagate Expansion 2TB"],
}
RU = {"Noutbuk": "Ноутбук", "Monitor": "Монитор", "Printer": "Принтер", "MFU": "МФУ", "Skaner": "Сканер",
      "Proyektor": "Проектор", "Kompyuter": "Компьютер", "Konditsioner": "Кондиционер", "UPS": "ИБП",
      "Router": "Роутер", "Televizor": "Телевизор", "Telefon apparati": "Телефон", "Stol": "Стол", "Kreslo": "Кресло",
      "Stul": "Стул", "Shkaf": "Шкаф", "Seyf": "Сейф", "Javon": "Стеллаж", "Tumba": "Тумба", "Doska": "Доска",
      "Kuler": "Кулер", "Isitgich": "Обогреватель", "Shreder": "Шредер", "Laminator": "Ламинатор",
      "Kalkulyator": "Калькулятор", "Veb-kamera": "Веб-камера", "Kolonka": "Колонки", "Planshet": "Планшет",
      "Server": "Сервер", "Kommutator": "Коммутатор", "Klaviatura": "Клавиатура", "Sichqoncha": "Мышь"}
PRICE = {"Noutbuk": 7_500_000, "Monitor": 2_100_000, "Printer": 2_600_000, "MFU": 3_900_000, "Skaner": 2_300_000,
         "Proyektor": 6_800_000, "Kompyuter": 8_200_000, "Konditsioner": 5_400_000, "Televizor": 5_100_000,
         "Server": 24_000_000, "Seyf": 3_200_000, "Muzlatgich": 3_600_000, "Fotoapparat": 6_900_000, "Planshet": 3_100_000}
# «O'xshash, lekin boshqa buyum» tuzoqlari: bittasi ikkala ro'yxatda, ikkinchisi faqat bittasida
FORCE = {("Noutbuk", "HP 250 G8"): "fuzzy", ("Noutbuk", "HP 250 G9"): "left",
         ("Printer", "Canon LBP6030"): "fuzzy", ("Printer", "Canon LBP2900"): "right",
         ("Monitor", "LG 24MK430"): "exact", ("Monitor", "LG 27MP400"): "right",
         ("UPS", "APC Back-UPS 650VA"): "exact", ("UPS", "APC Back-UPS 1100VA"): "left",
         ("Stol", "yozuv 140 sm"): "exact", ("Stol", "yozuv 160 sm"): "fuzzy"}
GROUPS = {"exact": 54, "fuzzy": 34, "left": 20, "right": 12}     # jami 120 xil buyum


def variant(cat, model, rnd):
    """Buxgalteriyada boshqacha yozilgan nom: kirill, tartib, qisqartma, probel."""
    ways = ["upper", "squeeze"]
    if cat in RU:
        ways += ["ru", "ru_tail", "ru", "ru_tail"]
    if cat in RU and len(cat) >= 6 and " " not in cat:
        ways.append("abbr")
    for _ in range(20):
        w = rnd.choice(ways)
        if w == "ru":
            s = f"{RU[cat]} {model}"
        elif w == "ru_tail":
            s = f"{model.replace(' G', 'G').replace('LBP', 'LBP-')} {RU[cat].lower()}"
        elif w == "upper":
            s = f"{cat.upper()}  {model}"
        elif w == "squeeze":
            s = f"{cat} {model.replace(' ', '').replace(',', '') if any(ch.isdigit() for ch in model) else model.replace(', ', ' ').upper()}"
        else:
            s = f"{cat[:4]}. {model}"
        if s != f"{cat} {model}":
            return s
    return f"{cat}  {model}"


def gen_inventory():
    rnd = random.Random(stable_seed("h01", "ombor"))
    items = [(c, m) for c, ms in CATALOG.items() for m in ms]
    assert len(items) == sum(GROUPS.values()), len(items)
    group = dict(FORCE)
    left = {g: n - sum(1 for v in FORCE.values() if v == g) for g, n in GROUPS.items()}
    rest = [it for it in items if it not in group]
    rnd.shuffle(rest)
    for g, n in left.items():
        for _ in range(n):
            group[rest.pop()] = g
    ombor, bux, pairs = [], [], []
    rooms = [f"{f}{n:02d}-xona" for f in (1, 2, 3) for n in range(1, 13)] + ["Ombor", "Majlislar zali", "Qabulxona"]
    for (cat, model) in items:
        g, name = group[(cat, model)], f"{cat} {model}"
        qty = rnd.choice([1, 1, 1, 2, 2, 3, 4, 6, 8, 12] if cat in ("Stol", "Stul", "Kreslo", "Monitor", "Klaviatura", "Sichqoncha") else [1, 1, 1, 2, 2, 3])
        price = int(PRICE.get(cat, 900_000) * rnd.uniform(0.7, 1.4) / 10_000) * 10_000
        if g in ("exact", "fuzzy", "left"):
            ombor.append({"name": name, "qty": qty, "room": rnd.choice(rooms), "state": "Ta'mirtalab" if rnd.random() < 0.08 else "Yaroqli"})
        if g in ("exact", "fuzzy", "right"):
            bname = variant(cat, model, rnd) if g == "fuzzy" else name
            bux.append({"name": bname, "qty": qty, "price": price,
                        "date": date(rnd.randint(2019, 2025), rnd.randint(1, 12), rnd.randint(1, 28))})
            if g != "right":
                pairs.append((name, bname, g))
    # soni mos kelmaydigan 6 ta buyum (ikkala ro'yxatda bor, lekin miqdori boshqa)
    both = [b for b in bux if any(p[1] == b["name"] and p[2] == "exact" for p in pairs)]
    for b in rnd.sample(both, 6):
        b["qty"] += rnd.choice([1, 1, 2])
        b["qty_diff"] = True
    for b in rnd.sample(bux, 5):                        # matn bo'lib qolgan raqam (payshanba darsi uchun)
        b["price_text"] = True
    rnd.shuffle(ombor)
    rnd.shuffle(bux)
    return ombor, bux, pairs, group


OMBOR, BUX, PAIRS, INV_GROUP = gen_inventory()
INV = dict(ombor=len(OMBOR), bux=len(BUX), exact=sum(1 for p in PAIRS if p[2] == "exact"),
           fuzzy=sum(1 for p in PAIRS if p[2] == "fuzzy"), left=GROUPS["left"], right=GROUPS["right"],
           qty_diff=sum(1 for b in BUX if b.get("qty_diff")), price_text=sum(1 for b in BUX if b.get("price_text")))
assert INV["ombor"] == INV["exact"] + INV["fuzzy"] + INV["left"] and INV["bux"] == INV["exact"] + INV["fuzzy"] + INV["right"], INV


def build_ombor(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ombor"
    write_table(ws, ["T/r", "Nomi", "Soni", "Joylashuvi", "Holati"],
                [[i, o["name"], o["qty"], o["room"], o["state"]] for i, o in enumerate(OMBOR, 1)],
                widths=[6, 44, 8, 18, 14], title="Ombor ro'yxati — inventarizatsiya (o'quv mashqi, o'ylab topilgan)")
    wb.save(path)


def build_bux(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Buxgalteriya"
    rows = [[i, b["name"], b["qty"], uzn(b["price"]) if b.get("price_text") else b["price"], b["date"]]
            for i, b in enumerate(BUX, 1)]
    write_table(ws, ["T/r", "Asosiy vosita nomi", "Soni", "Boshlang'ich qiymati (so'm)", "Hisobga olingan sana"], rows,
                widths=[6, 44, 8, 22, 16], title="Buxgalteriya hisobi — asosiy vositalar (o'quv mashqi, o'ylab topilgan)",
                date_cols=(5,), money_cols=(4,))
    wb.save(path)


F_OMBOR = W.file("ombor.xlsx", build_ombor)
F_BUX = W.file("buxgalteriya.xlsx", build_bux)


# ══════════════════════════════════════════════════════════════════ 4–5-fayl: ko'chirma va byudjet shabloni
# (toifa, tavsif, oyiga necha marta, summa oralig'i)
SPEND = [("Oziq-ovqat", ["SUPERMARKET 12", "SUPERMARKET CHILONZOR", "BOZOR SAVDO", "NON-SUT MAHSULOTLARI"], 13, (45_000, 380_000)),
         ("Transport", ["TAKSI", "AZS YOQILG'I", "TRANSPORT KARTASI"], 9, (12_000, 210_000)),
         ("Kafe va ovqatlanish", ["KAFE", "OSHXONA TUSHLIK", "FAST-FUD"], 6, (28_000, 190_000)),
         ("Sog'liq", ["DORIXONA", "KLINIKA"], 2, (35_000, 260_000)),
         ("Kiyim va uy", ["KIYIM DO'KONI", "XO'JALIK MOLLARI", "ONLAYN DO'KON"], 2, (90_000, 640_000)),
         ("Boshqa", ["P2P O'TKAZMA", "KITOB DO'KONI", "SOVG'A"], 2, (50_000, 400_000))]
# (toifa, tavsif, kun, summa) — har oy bir xil summada takrorlanadi
FIXED = [("Kommunal", "ELEKTR ENERGIYA", 7, None), ("Kommunal", "TABIIY GAZ", 7, None),
         ("Kommunal", "ICHIMLIK SUVI", 8, 48_000), ("Kommunal", "CHIQINDI OLIB KETISH", 8, 30_000),
         ("Aloqa", "UY INTERNETI", 3, 149_000), ("Aloqa", "MOBIL ALOQA TARIFI", 15, 70_000),
         ("Obuna va abonement", "ONLAYN KINOTEATR OBUNA", 11, 39_000), ("Obuna va abonement", "MUSIQA OBUNA", 11, 19_000),
         ("Obuna va abonement", "BULUTLI XOTIRA OBUNA", 19, 15_000), ("Obuna va abonement", "SPORT ZALI ABONEMENT", 2, 350_000),
         ("Obuna va abonement", "BOLALAR TO'GARAGI", 5, 400_000)]
ST_MONTHS = [(2026, 7), (2026, 8), (2026, 9)]


def gen_statement():
    rnd = random.Random(stable_seed("h01", "kochirma"))
    rows = []
    for (y, m) in ST_MONTHS:
        last = (date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)).day
        rows.append({"date": date(y, m, 5), "desc": "ISH HAQI", "inc": 6_200_000, "cat": None})
        rows.append({"date": date(y, m, 20), "desc": "ISH HAQI (AVANS)", "inc": 3_800_000, "cat": None})
        for cat, desc, day, amt in FIXED:
            a = amt if amt else int(rnd.uniform(140_000, 330_000) / 1000) * 1000
            rows.append({"date": date(y, m, day), "desc": desc, "out": a, "cat": cat, "fixed": amt is not None})
        for cat, descs, n, (lo, hi) in SPEND:
            for _ in range(n + rnd.choice([-1, 0, 0, 1])):
                rows.append({"date": date(y, m, rnd.randint(1, last)), "desc": rnd.choice(descs),
                             "out": int(rnd.triangular(lo, hi, lo + (hi - lo) * 0.3) / 1000) * 1000, "cat": cat})
    rows.sort(key=lambda r: (r["date"], r["desc"]))
    return rows


STATEMENT = gen_statement()
_n = len(ST_MONTHS)
CAT_AVG = Counter()
for r in STATEMENT:
    if r.get("out"):
        CAT_AVG[r["cat"]] += r["out"] / _n
CAT_AVG = {k: int(round(v / 1000)) * 1000 for k, v in CAT_AVG.items()}
REC = [(d, a) for c, d, day, a in FIXED if a]                      # har oy bir xil summa
REC_MONTH = sum(a for _, a in REC)
REC_YEAR = REC_MONTH * 12
SUBS_MONTH = sum(a for c, d, day, a in FIXED if c == "Obuna va abonement")
BUDGET_CATS = ["Oziq-ovqat", "Kommunal", "Transport", "Aloqa", "Obuna va abonement", "Kafe va ovqatlanish",
               "Sog'liq", "Kiyim va uy", "Boshqa"]
assert set(BUDGET_CATS) == set(CAT_AVG), set(CAT_AVG) ^ set(BUDGET_CATS)


def build_statement(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ko'chirma"
    write_table(ws, ["Sana", "Tavsif", "Kirim (so'm)", "Chiqim (so'm)"],
                [[r["date"], r["desc"], r.get("inc"), r.get("out")] for r in STATEMENT],
                widths=[13, 34, 16, 16], title="Karta bo'yicha ko'chirma, iyul–sentabr 2026 (o'quv mashqi, o'ylab topilgan)",
                date_cols=(1,), money_cols=(3, 4))
    wb.save(path)


def build_budget(path):
    """Shablon: sariq kataklar to'ldiriladi, qolgani formula. Namuna qiymatlar — o'ylab topilgan ko'chirmadan."""
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Byudjet"
    f, b = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    yellow, grey = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="E7E6E6")
    ws["A1"], ws["A1"].font = "Oilaviy byudjet — oylik reja va haqiqat", Font(name="Arial", size=12, bold=True)
    ws["A2"] = "Sariq kataklarni o'zingiz to'ldiring; «Farq» va «Bajarilishi» o'zi hisoblanadi. Hozirgi raqamlar — namuna (o'ylab topilgan ko'chirmadan)."
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="7F7F7F")
    for c, h in enumerate(["Toifa", "Reja (so'm)", "Haqiqat (so'm)", "Farq (so'm)", "Bajarilishi"], 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font, cell.fill, cell.alignment = b, grey, Alignment(horizontal="center")
    r0 = 5
    for i, cat in enumerate(BUDGET_CATS):
        r = r0 + i
        fact = CAT_AVG[cat]
        plan = int(round(fact / 100_000)) * 100_000 or 100_000       # namuna reja: yaxlit summa
        ws.cell(row=r, column=1, value=cat).font = f
        for c, v in ((2, plan), (3, fact)):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.fill, cell.number_format = f, yellow, "#,##0"
        ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = '#,##0;[Red]-#,##0'
        ws.cell(row=r, column=5, value=f'=IF(B{r}=0,"",C{r}/B{r})').number_format = "0%"
        ws.cell(row=r, column=4).font = ws.cell(row=r, column=5).font = f
    rt, re_ = r0 + len(BUDGET_CATS), r0 + len(BUDGET_CATS) - 1
    ws.cell(row=rt, column=1, value="Jami").font = b
    for c in (2, 3, 4):
        L = get_column_letter(c)
        cell = ws.cell(row=rt, column=c, value=f"=SUM({L}{r0}:{L}{re_})")
        cell.font, cell.number_format = b, '#,##0;[Red]-#,##0'
    ws.cell(row=rt, column=5, value=f'=IF(B{rt}=0,"",C{rt}/B{rt})').number_format = "0%"
    ws.cell(row=rt, column=5).font = b
    ws.cell(row=rt + 2, column=1, value="Oylik daromad (so'm)").font = f
    inc = ws.cell(row=rt + 2, column=2, value=10_000_000)
    inc.font, inc.fill, inc.number_format = f, yellow, "#,##0"
    ws.cell(row=rt + 3, column=1, value="Daromaddan qoladi (so'm)").font = b
    left = ws.cell(row=rt + 3, column=2, value=f"=B{rt + 2}-C{rt}")
    left.font, left.number_format = b, '#,##0;[Red]-#,##0'
    for c, w in enumerate([30, 16, 16, 16, 13], 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    ws2 = wb.create_sheet("Takroriy to'lovlar")
    ws2["A1"], ws2["A1"].font = "Har oy takrorlanadigan to'lovlar — yillik summada", Font(name="Arial", size=12, bold=True)
    ws2["A2"] = "Sariq kataklarga o'z to'lovlaringizni yozing. «Yiliga» ustuni o'zi hisoblanadi."
    ws2["A2"].font = Font(name="Arial", size=9, italic=True, color="7F7F7F")
    for c, h in enumerate(["To'lov", "Oyiga (so'm)", "Yiliga (so'm)"], 1):
        cell = ws2.cell(row=4, column=c, value=h)
        cell.font, cell.fill, cell.alignment = b, grey, Alignment(horizontal="center")
    names = [d.capitalize() for d, _ in REC] + [None] * 5       # 5 ta bo'sh qator — o'zingiznikini qo'shish uchun
    for i, name in enumerate(names):
        r = 5 + i
        a = ws2.cell(row=r, column=1, value=name)
        v = ws2.cell(row=r, column=2, value=REC[i][1] if i < len(REC) else None)
        a.font, v.font, a.fill, v.fill, v.number_format = f, f, yellow, yellow, "#,##0"
        y = ws2.cell(row=r, column=3, value=f'=IF(B{r}="","",B{r}*12)')
        y.font, y.number_format = f, "#,##0"
    rt2 = 5 + len(names)
    ws2.cell(row=rt2, column=1, value="Jami").font = b
    for c in (2, 3):
        L = get_column_letter(c)
        cell = ws2.cell(row=rt2, column=c, value=f"=SUM({L}5:{L}{rt2 - 1})")
        cell.font, cell.number_format = b, "#,##0"
    for c, w in enumerate([34, 16, 16], 1):
        ws2.column_dimensions[get_column_letter(c)].width = w
    wb.save(path)


F_STATEMENT = W.file("kochirma_namuna.xlsx", build_statement)
F_BUDGET = W.file("byudjet_shablon.xlsx", build_budget)


# ══════════════════════════════════════════════════════════════════ postlar
def pc(x):
    return f"{x:.0f} %"


# ───────────────────────────── DUSHANBA · Hafta sinovi
_marked = {i for i, r in enumerate(APPEALS)
           if r.get("dup") or r.get("dup_of") or r.get("spike") or (r["dept"] == WORST_DEPT and r.get("late"))}
W.post("dush", rubric="#sinov", emoji="🔍", title="Sinov: 500 qator, 3 yashirin muammo",
       head="500 qator, 3 ta yashirin muammo",
       body=f"""Rahbar so'radi: «Murojaatlar bilan ishlashda ahvol qanday?» Qo'lingizda bir yillik reestr — {A['total']} qator.

📎 Keyingi xabarda <b>murojaatlar_2025.xlsx</b> — o'ylab topilgan idoraning reestri. Ichiga ataylab 3 ta muammo yashirilgan:

1️⃣ Bitta bo'lim muddatni boshqalardan bir necha barobar ko'p buzadi. Umumiy foizga qarasangiz — hammasi joyida.
2️⃣ Bir oyda murojaatlar keskin ko'paygan, sababi — bitta mavzu.
3️⃣ Ayrim murojaatlar ikki marta ro'yxatga olingan.

✍️ <b>Vazifa:</b> Excelda qo'lda topib ko'ring va vaqtni belgilang: qaysi bo'lim, qaysi oy, nechta takror? Javobingizni izohda yozing.

Ertaga shu ishni AI bilan qilamiz va javoblarni solishtiramiz.""",
       new="AI ga jadvalni matn qilib ko'chirish shart emas. Faylning o'zini bersangiz, u faylni ochib, ichida ishlaydi.",
       take="mashq fayli <b>murojaatlar_2025.xlsx</b>. Bu haftaning hamma darsi shu faylda.",
       card={"kind": "stats", "title": "Qo'lda topa olasizmi?",
             "stats": [(str(A["total"]), "qator"), ("3", "yashirin muammo", True)],
             "dots": {"n": A["total"], "cols": 40, "marked": _marked},
             "caption": "har nuqta — bitta qator · murojaatlar_2025.xlsx", "take": "Mashq fayli: bir yillik murojaatlar reestri"},
       attach=[{"file": F_APPEALS, "caption": "📎 <b>Hafta mashq fayli.</b> O'ylab topilgan ma'lumot — haqiqiy idora yoki shaxslarga aloqasi yo'q.\n\n{{CHANNEL}}"}],
       notes=(f"Fayldagi javoblar (tozalangan {A['uniq']} qator bo'yicha): {A['worst']} — muddatida {pc(A['worst_pct'])}, "
              f"qolganlar {pc(A['others_pct'])}, umumiy {pc(A['overall'])}; {OYLAR[A['month'] - 1]} — {A['month_n']} ta murojaat "
              f"(boshqa oylarda o'rtacha {A['month_avg']:.0f}), shundan {A['month_topic'][1]} tasi «{A['month_topic'][0]}»; "
              f"takror — {A['dups']} ta. Javoblar seshanba postida spoyler ostida."),
       scene="a tall stack of paper registry sheets with a magnifying glass over it; three small accent-coloured flags stick out of the stack at different heights.")

# ───────────────────────────── SESHANBA · Dars
W.post("sesh", rubric="#dars", emoji="📘", title="Dars: «hisobla» emas, «kod bilan hisobla»",
       head="AI hisobda adashadi. Kod yozdirsangiz — adashmaydi",
       body=f"""Jadvalni AI ga matn qilib tashlab «jami qancha?» desangiz, javob ko'pincha xato chiqadi. Sababi: til modeli raqamni hisoblamaydi, keyingi so'zni taxmin qiladi. {A['total']} qatorni u «ko'z bilan» qo'sha olmaydi.

<b>Yechim:</b> faylning o'zini yuklang va hisobni kod bilan qildiring. Shunda AI faylni ochadi, dastur yozadi va ishga tushiradi — hisobni model emas, kompyuter bajaradi. Bu imkoniyat ChatGPT, Claude va Gemini'da bor. Bepul tarifda ham sinab ko'rish mumkin, lekin kunlik limit bor.

📎 Fayl: kechagi <b>murojaatlar_2025.xlsx</b>

<b>1-qadam. Avval faylni tekshirtiring</b>
{pre('''Bu faylni kod bilan och. Hech narsa hisoblama. Avval ayt:
nechta qator va ustun bor, sarlavha qaysi qatorda,
har ustunda nima bor, qayerda bo'sh katak, takror qator
yoki sana matn bo'lib qolgan.''')}
<i>Nega:</i> iflos ma'lumotdan chiqqan xulosa ham xato. Odam buni bir soatda topadi, kod — 10 soniyada.

<b>2-qadam. Tozalash qoidasini o'zingiz belgilang</b>
{pre('''Takror murojaatlarni olib tashla: murojaatchi kodi, sana va mavzu
bir xil bo'lsa — bu takror. "Toshkent sh.", "г.Ташкент" kabi
yozuvlarni bitta ko'rinishga keltir.
Nimani o'zgartirganingni ro'yxat qilib ber.''')}
<i>Nega:</i> «takror» nima ekanini siz bilasiz, AI emas.

<b>3-qadam. Savol bering</b>
{pre('''Kod bilan hisobla: har bo'lim bo'yicha murojaatlar soni,
muddatida bajarilganlar foizi, o'rtacha ko'rib chiqish kuni.
Jadval qilib ber, eng yomon 3 bo'limni ajrat.''')}

<b>4-qadam. So'ramagan narsangizni so'rang</b>
{pre('''Bu ma'lumotda rahbar bilishi kerak bo'lgan, lekin men so'ramagan
3 ta g'ayrioddiy holatni top. Har biriga raqamli dalil keltir.''')}
<i>Nega:</i> jadvalni Excel ham tuzadi. AI ning foydasi — siz qidirishni o'ylamagan narsani topishida.

<b>5-qadam. Tayyor fayl oling</b>
{pre('''Natijani Excel fayl qilib ber: 1-varaq — tozalangan ma'lumot,
2-varaq — bo'limlar jadvali, 3-varaq — grafik.
Alohida yarim betlik ma'lumotnoma matnini yoz.''')}

✅ <b>Tekshirish (1 daqiqa):</b> AI dan jami qatorlar sonini so'rang va Excel pastidagi hisoblagich bilan solishtiring. Mos kelmasa — «kodni ko'rsat» deng.

🔑 <b>Kechagi sinov javoblari</b> (bosing):
<tg-spoiler>1) {A['worst']}: muddatida {pc(A['worst_pct'])}, qolgan bo'limlarda {pc(A['others_pct'])}. Umumiy ko'rsatkich {pc(A['overall'])} — shuning uchun ko'rinmaydi.
2) {OYLAR[A['month'] - 1].capitalize()}: {A['month_n']} ta murojaat, boshqa oylarda o'rtacha {A['month_avg']:.0f} ta. Shundan {A['month_topic'][1]} tasi — «{A['month_topic'][0]}».
3) Takror ro'yxatga olingan: {A['dups']} ta.</tg-spoiler>

⚠️ Xizmat hujjatlari va shaxsiy ma'lumotlar AI ga yuklanmaydi. Mashq — faqat o'ylab topilgan faylda.""",
       new="«hisobla» emas, «<b>kod bilan</b> hisobla». Shu ikki so'z javobni taxmindan hisobga aylantiradi.",
       take="5 so'rovli tahlil zanjiri — istalgan jadvalga mos. Postni saqlab qo'ying.",
       card={"kind": "chain", "title": "«Hisobla» emas, «kod bilan hisobla»",
             "steps": [("Faylni tekshirtiring", "bo'sh katak · takror · matn-sana"),
                       ("Tozalash qoidasini bering", "nimani o'zgartirganini yozsin"),
                       ("Savol bering", "«kod bilan hisobla»"),
                       ("So'ramaganingizni so'rang", "3 ta g'ayrioddiy holat, dalil bilan"),
                       ("Tayyor fayl oling", "Excel + yarim betlik ma'lumotnoma")],
             "take": "5 so'rovli tahlil zanjiri"},
       notes=(f"Vositalar holati {CHECKED} da rasmiy yordam sahifalaridan tekshirildi (manbalarga qarang). Gemini bepul tarifida "
              "kod bilan tahlil qilish rasmiy sahifada aniq yozilmagan — tasdiqlashdan oldin o'zingiz bir sinab ko'ring. "
              f"Faylda ataylab qo'yilgan: {A['dirty_regions']} ta turlicha yozilgan hudud, {A['blanks']} ta bo'sh «Mas'ul bo'lim», "
              f"{A['text_dates']} ta matn bo'lib qolgan sana, sarlavha 3-qatorda."),
       sources=SRC_TOOLS,
       scene="a spreadsheet sheet feeding into a small tidy machine with five numbered gears in a row; out of the machine comes a clean one-page report.")

# ───────────────────────────── CHORSHANBA · Ustalik
_fz = [p for p in PAIRS if p[2] == "fuzzy"]
_ex1 = next(p for p in _fz if p[0] == "Noutbuk HP 250 G8")
_ex2 = next(p for p in _fz if p[0] == "Printer Canon LBP6030")
W.post("chor", rubric="#ustalik", emoji="🔗", title="Ustalik: ikki ro'yxatni solishtirish",
       head="Ikki ro'yxatni solishtirish",
       body=f"""Har idorada bor ish: ombor ro'yxati buxgalteriya hisobi bilan mos kelishi kerak. Qo'lda — yarim kun. Excel formulasi esa «{_ex1[0]}» bilan «{_ex1[1]}» bitta buyum ekanini tanimaydi.

📎 Keyingi xabarlarda: <b>ombor.xlsx</b> ({INV['ombor']} qator) va <b>buxgalteriya.xlsx</b> ({INV['bux']} qator), ikkalasi o'ylab topilgan.

<b>1. Yozilishini ko'rsin</b>
{pre('''Ikki faylni kod bilan och. Har birida nomlar qanday yozilganini
ko'rsat: lotin yoki kirill, qisqartmalar, ortiqcha probel va belgilar.''')}

<b>2. Bir xil ko'rinishga keltirsin</b>
{pre('''Nomlarni bir xil ko'rinishga keltir. Asl ustunga tegma,
yonida yangi ustun och.''')}

<b>3. Uch guruhga ajratsin</b>
{pre('''Ikki ro'yxatni solishtir va uch guruhga ajrat:
aniq mos keldi / o'xshash, lekin ishonch yo'q / faqat bittasida bor.
"O'xshash" guruhida ikkala yozuvni yonma-yon ko'rsat.
Mos kelganlarda soni farq qilsa — alohida belgilab ber.''')}

<b>4. Faylga chiqarsin</b>
{pre('''Natijani Excel qilib ber, har guruh alohida varaqda.''')}

<i>Nega uch guruh:</i> xavfli joy — «o'xshash» guruhi. Uni siz hal qilasiz. «O'zing hal qil» desangiz, AI HP 250 G8 bilan HP 250 G9 ni bitta buyum qilib yuboradi.

🔑 <b>O'zingizni tekshiring</b> (bosing):
<tg-spoiler>Aniq mos: {INV['exact']} ta · boshqacha yozilgan, lekin bitta buyum: {INV['fuzzy']} ta · faqat omborda: {INV['left']} ta · faqat buxgalteriyada: {INV['right']} ta · soni farq qiladi: {INV['qty_diff']} ta.</tg-spoiler>

🔒 Ro'yxatda shaxsiy yoki xizmat ma'lumoti bo'lsa — faylni yuklamang. Ustun nomlarini va 2–3 ta o'ylab topilgan qatorni bering, AI dan Excel formulasi yoki makros so'rang: hisob o'z kompyuteringizda bajariladi.""",
       new="AI dan javob emas, <b>saralash</b> so'rang: aniq / shubhali / yo'q. Shubhalisini o'zingiz hal qilasiz.",
       take="4 so'rov va ikki mashq fayli — o'z ro'yxatlaringizga ham shu tartib ishlaydi.",
       card={"kind": "pairs", "title": "Ikki ro'yxatni solishtirish", "heads": ("OMBOR", "BUXGALTERIYA"),
             "left": [_ex1[0], _ex2[0], "Stol yozuv 140 sm", "Monitor LG 24MK430", "Noutbuk HP 250 G9", "UPS APC Back-UPS 1100VA"],
             "right": [_ex1[1], "Stol yozuv 140 sm", _ex2[1], "Monitor LG 24MK430", "Printer Canon LBP2900", "Monitor LG 27MP400"],
             "links": [(0, 0, "~"), (1, 2, "~"), (2, 1, "="), (3, 3, "=")],
             "take": "4 so'rov va ikki mashq fayli"},
       attach=[{"file": F_OMBOR, "caption": "📎 Ombor ro'yxati (o'ylab topilgan)."},
               {"file": F_BUX, "caption": "📎 Buxgalteriya hisobi (o'ylab topilgan).\n\n{{CHANNEL}}"}],
       notes=(f"Fayllar: jami {sum(GROUPS.values())} xil buyum. Tuzoqlar: HP 250 G8 (ikkalasida, boshqacha yozilgan) va HP 250 G9 (faqat omborda); "
              "Canon LBP6030 va LBP2900; LG 24MK430 va 27MP400; APC 650VA va 1100VA. "
              f"Buxgalteriyada {INV['price_text']} ta qiymat matn bo'lib yozilgan — ertangi #ehtiyot posti uchun."),
       scene="two paper lists side by side; thin threads connect matching lines between them — some threads solid, some dashed, a few lines left unconnected.")

# ───────────────────────────── PAYSHANBA · Ehtiyot bo'ling
W.post("pay", rubric="#ehtiyot", emoji="⚠️", title="Ehtiyot: fayl tahlili adashadigan 5 joy",
       head="Fayl tahlili adashadigan 5 joy",
       body=f"""Kod bilan hisoblash taxmindan yaxshi. Lekin kod ham faylni noto'g'ri o'qishi mumkin. Beshta joy va har biriga bitta nazorat savoli:

<b>1. Sana.</b> 03.04.2025 — uchinchi aprelmi, to'rtinchi martmi? Kod ko'pincha amerikacha o'qiydi.
→ <code>Sanalar kun.oy.yil formatida. Eng erta va eng kech sanani ko'rsat.</code>

<b>2. Matn bo'lib qolgan raqam.</b> «1 250 000» probel bilan yozilsa, yig'indiga kirmay qoladi.
→ <code>Qaysi ustunlarda raqam matn bo'lib turibdi? Nechta katak?</code>

<b>3. Sarlavha.</b> Jadval tepasida nom yoki birlashtirilgan katak bo'lsa, ustunlar siljib ketadi.
→ <code>Sarlavhani qaysi qatordan olding? Ustun nomlarini sanab ber.</code>

<b>4. Faqat boshini ko'rish.</b> AI birinchi qatorlarga qarab xulosa aytishi mumkin.
→ <code>Nechta qator qayta ishlandi?</code> Javob fayldagi qatorlar soniga teng bo'lishi kerak.

<b>5. Kodsiz javob.</b> AI hisoblamasdan raqam aytadi.
→ <code>Kodni va uning natijasini ko'rsat.</code>

🧪 <b>Sinab ko'ring:</b> bu hafta fayllarida uchala tuzoq ham bor. murojaatlar_2025.xlsx da {A['text_dates']} ta sana matn bo'lib yozilgan va sarlavha 3-qatorda, buxgalteriya.xlsx da {INV['price_text']} ta qiymat matn. AI ularni o'zi aytdimi?""",
       new="AI ning javobini emas, <b>hisoblash yo'lini</b> tekshirasiz. Besh savol — ikki daqiqa.",
       take="5 nazorat savoli — rasmni saqlab qo'ying, har tahlildan keyin bering.",
       card={"kind": "checks", "title": "Fayl tahlili adashadigan 5 joy",
             "items": [("Sana: 03.04 — aprelmi, martmi?", "Eng erta va eng kech sanani ko'rsat"),
                       ("Matn bo'lib qolgan raqam", "Qaysi ustunda raqam matn bo'lib turibdi?"),
                       ("Sarlavha siljishi", "Sarlavhani qaysi qatordan olding?"),
                       ("Faqat boshini ko'rish", "Nechta qator qayta ishlandi?"),
                       ("Kodsiz javob", "Kodni va natijasini ko'rsat")],
             "take": "5 nazorat savoli — har tahlildan keyin"},
       notes="Rasmning o'zi — «olib keting» kartochkasi: 5 tuzoq va 5 savol. Savollar postda <code> ichida — bosilsa nusxalanadi.",
       scene="a paper checklist with five rows on a clipboard; a small warning triangle stamp in the accent colour on the top corner.")

# ───────────────────────────── JUMA · Bonus
_bars = [(c, CAT_AVG[c], c == "Obuna va abonement") for c in sorted(BUDGET_CATS, key=lambda c: -CAT_AVG[c])]
W.post("juma", rubric="#bonus", emoji="💳", title="Bonus: pul qayerga ketyapti",
       head="Pul qayerga ketyapti",
       body=f"""Bu hafta o'rganganingiz uyda ham ishlaydi: bank ko'chirmasi — xuddi o'sha jadval.

📎 Avval namunada mashq qiling: keyingi xabarda <b>kochirma_namuna.xlsx</b> — o'ylab topilgan 3 oylik ko'chirma.

<b>1. Toifalarga ajratsin</b>
{pre('''Bu ko'chirmani kod bilan och. Har chiqimni toifaga ajrat:
oziq-ovqat, transport, kommunal, aloqa, obuna, kafe, sog'liq, kiyim, boshqa.
Ishonching komil bo'lmagan to'lovlarni alohida ko'rsat.''')}

<b>2. Takroriy to'lovlarni topsin</b>
{pre('''Har oy bir xil summada takrorlanadigan to'lovlarni top.
Har birining yillik summasini hisobla.''')}

<b>3. Byudjet tuzsin</b>
{pre('''Toifalar bo'yicha oylik o'rtacha chiqim jadvalini tuz.
Qaysi uch toifada tejash imkoni borligini raqam bilan ko'rsat.''')}

🔑 <b>Namunadagi javob</b> (bosing):
<tg-spoiler>Har oy bir xil summada takrorlanadigan {len(REC)} ta to'lov: oyiga {uzn(REC_MONTH)} so'm — yiliga {uzn(REC_YEAR)} so'm.</tg-spoiler>

Keyin o'z ko'chirmangizda sinang. Bank ilovasidan Excel qilib oling va yuklashdan oldin <b>karta raqami, F.I.Sh. va hisob raqami</b> bor qatorlarni o'chiring.""",
       new="mayda takroriy to'lov oyda sezilmaydi — yillik summada ko'rinadi.",
       take="<b>byudjet_shablon.xlsx</b> — reja, haqiqat va takroriy to'lovlarning yillik hisobi. Sariq kataklarni o'zingiz to'ldirasiz.",
       card={"kind": "bars", "title": "Pul qayerga ketyapti", "sub": "Namunaviy oila, oylik o'rtacha chiqim (so'm)",
             "items": _bars,
             "note": (f"yiliga {uzn(REC_YEAR)} so'm", f"{len(REC)} ta takroriy to'lov: oyiga {uzn(REC_MONTH)} so'm"),
             "take": "Byudjet shabloni va namunaviy ko'chirma"},
       attach=[{"file": F_STATEMENT, "caption": "📎 Namunaviy ko'chirma (o'ylab topilgan) — avval shunda mashq qiling."},
               {"file": F_BUDGET, "caption": "📎 Byudjet shabloni: sariq kataklarni to'ldirasiz, qolgani o'zi hisoblanadi.\n\n{{CHANNEL}}"}],
       notes=("Ko'chirma va raqamlar o'ylab topilgan. Takroriy to'lovlar: " + ", ".join(f"{d.lower()} {uzn(a)}" for d, a in REC)
              + f". Shundan obuna va abonementlar oyiga {uzn(SUBS_MONTH)} so'm."),
       scene="a household budget sheet with a few coins and small receipt slips; one repeating small receipt is highlighted in the accent colour and multiplied into a tall stack.")

W.build()

if "--cards" not in " ".join(W.args):
    print(f"Javoblar: {A['worst']} {A['worst_pct']:.1f}% (qolganlar {A['others_pct']:.1f}%, umumiy {A['overall']:.1f}%, farq {A['ratio']:.1f} barobar); "
          f"{OYLAR[A['month'] - 1]} {A['month_n']} (o'rtacha {A['month_avg']:.1f}); takror {A['dups']}.")
    print(f"Ro'yxatlar: ombor {INV['ombor']}, buxgalteriya {INV['bux']}, aniq {INV['exact']}, o'xshash {INV['fuzzy']}, "
          f"faqat ombor {INV['left']}, faqat bux {INV['right']}, soni farq {INV['qty_diff']}.")
    print(f"Takroriy to'lovlar: oyiga {uzn(REC_MONTH)}, yiliga {uzn(REC_YEAR)}.")
